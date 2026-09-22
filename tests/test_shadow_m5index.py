import csv
import hashlib
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHADOW = ROOT / "shadow-m5index"
DATA = SHADOW / "data"
CSV_PATH = DATA / "assets" / "SHADOW_M5Index_Assets_Index_v0.1.csv"
SOURCE_WORKBOOK = DATA / "source" / "SHADOW_M5Index_Master_Dataset_v0.1.source.xlsx"
MASTER_WORKBOOK = DATA / "SHADOW_M5Index_Master_Dataset_v0.1.xlsx"
HANDOFF = SHADOW / "reference" / "SHADOW-M5INDEX-CANONICAL-ENGINEERING-HANDOFF.md"
GENERATOR = SHADOW / "tools" / "build_master_dataset.py"

EXPECTED_SHA256 = {
    SOURCE_WORKBOOK: "11a4fce375128f2ffb6e6833a3da916614e52587eb20afae78c56835ccdcb148",
    CSV_PATH: "65fab1a869b1a0eb3d59d06a7e0f1dadbbb2f96752b33210990ed8eed984efe4",
    HANDOFF: "2d572b4d5325a581052890f16be30700e2884e7a5879b96193bff520d9197a62",
}

REQUIRED_SHEETS = [
    "README",
    "Assets",
    "Ownership_Capital",
    "Debt_Liens",
    "TitleChain_Asset_State",
    "Rights_Encumbrances",
    "Claims_Obligations",
    "Environmental_State",
    "Dispositions_Deal_Flow",
    "SHADOW_Scores",
    "SHADOW_CAMEL",
    "Orbitalys_Threats",
    "Peoples_Trust_Review",
    "Evidence_Sources",
    "Research_Queue_T1",
    "Research_Queue_T2",
    "Research_Queue_T3",
]

EVIDENCE_STATES = {
    "SUBMITTED",
    "SOURCE_VERIFIED",
    "CORROBORATED",
    "CANONICAL",
    "DISPUTED",
    "UNRESOLVED",
    "REJECTED",
    "CORRECTED",
    "SUPERSEDED",
}

REQUIRED_DIRECTORIES = {
    "methodology",
    "schemas",
    "data/assets",
    "data/institutions",
    "data/dispositions",
    "data/rights",
    "data/obligations",
    "data/jurisdictions",
    "data/sources",
    "reports",
    "examples",
    "visuals",
    "reference/graph",
    "reference/scoring",
    "reference/titlechain-adapter",
    "reference/orbitalys-adapter",
    "reference/camel",
    "reference/people-trust-review",
    "tests",
}

SHADOW_ISSUE_FORMS = [
    ROOT / ".github" / "ISSUE_TEMPLATE" / "shadow-new-asset.yml",
    ROOT / ".github" / "ISSUE_TEMPLATE" / "shadow-add-evidence.yml",
    ROOT / ".github" / "ISSUE_TEMPLATE" / "shadow-correction.yml",
]

REPORT_PAGE = SHADOW / "reports" / "SHADOW-CAMEL-REPORT-Q3-2026.md"
REPORT_PDF = SHADOW / "reports" / "M5Index_SHADOW_CAMEL_Report_Q3_2026.pdf"
REPORT_DOCX = SHADOW / "reports" / "M5Index_SHADOW_CAMEL_Report_Q3_2026.docx"
REPORT_COVER = SHADOW / "reports" / "shadow-camel-report-q3-2026-cover.png"


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return headers, [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]


def test_imported_source_fingerprints_are_immutable():
    assert {path: sha256(path) for path in EXPECTED_SHA256} == EXPECTED_SHA256


def test_pr1_module_structure_is_present():
    assert all((SHADOW / path).is_dir() for path in REQUIRED_DIRECTORIES)


def workbook_values(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    return {
        sheet.title: [
            [cell.value for cell in row]
            for row in sheet.iter_rows()
        ]
        for sheet in workbook.worksheets
    }


def test_normalized_workbook_rebuild_is_content_reproducible(tmp_path):
    rebuilt = tmp_path / MASTER_WORKBOOK.name
    subprocess.run(
        [
            sys.executable,
            str(GENERATOR),
            "--source-workbook",
            str(SOURCE_WORKBOOK),
            "--source-assets-csv",
            str(CSV_PATH),
            "--output",
            str(rebuilt),
        ],
        check=True,
    )
    assert workbook_values(rebuilt) == workbook_values(MASTER_WORKBOOK)


def test_asset_csv_has_unique_stable_ids_and_expected_tranches():
    with CSV_PATH.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))

    ids = [row["Asset_ID"] for row in rows]
    assert len(rows) == len(ids) == len(set(ids)) == 36
    assert ids == [f"SHD-{number:03d}" for number in range(1, 37)]
    assert {row["Tranche"] for row in rows} == {"Tranche 1", "Tranche 2", "Tranche 3"}
    assert {row["Evidence_State"] for row in rows} <= EVIDENCE_STATES
    assert all(row["Evidence_State"] != "CANONICAL" for row in rows)


def test_normalized_workbook_has_required_public_surfaces_and_no_external_links():
    workbook = load_workbook(MASTER_WORKBOOK, read_only=False, data_only=False)
    assert workbook.sheetnames == REQUIRED_SHEETS
    assert not workbook._external_links
    assert not workbook.defined_names

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert not re.match(r"^\s*[=+@]", cell.value)


def test_normalized_assets_match_csv_and_evidence_never_promotes():
    workbook = load_workbook(MASTER_WORKBOOK, read_only=True, data_only=False)
    _, assets = table(workbook["Assets"])
    with CSV_PATH.open(newline="", encoding="utf-8-sig") as handle:
        csv_rows = list(csv.DictReader(handle))

    assert [row["Asset_ID"] for row in assets] == [row["Asset_ID"] for row in csv_rows]
    assert all(row["Evidence_State"] in EVIDENCE_STATES for row in assets)
    assert all(row["Evidence_State"] != "CANONICAL" for row in assets)


def test_scores_camel_and_peoples_trust_remain_non_authoritative():
    workbook = load_workbook(MASTER_WORKBOOK, read_only=True, data_only=False)
    _, scores = table(workbook["SHADOW_Scores"])
    camel_headers, camel_rows = table(workbook["SHADOW_CAMEL"])
    _, reviews = table(workbook["Peoples_Trust_Review"])

    assert len(scores) == 36
    assert all(row["Score_Status"] == "NOT SCORED" for row in scores)
    assert "Status" in camel_headers
    assert camel_rows == []
    assert all(row["Decision"] == "NEED MORE EVIDENCE" for row in reviews)

    text = (SHADOW / "README.md").read_text()
    assert "not an official CAMELS rating" in text
    assert "nomination creates no right" in text


def test_spring_street_is_a_hold_research_lead_not_an_acquisition_claim():
    workbook = load_workbook(MASTER_WORKBOOK, read_only=True, data_only=False)
    _, assets = table(workbook["Assets"])
    spring = next(row for row in assets if row["Asset_ID"] == "SHD-029")

    assert spring["Disposition_State"] == "P2 MARKETED"
    assert spring["Transaction_Proof_State"] == "S0 LISTED"
    assert spring["Ownership_Resolution_State"] == "O0 BUYER_UNDISCLOSED"
    assert spring["Environmental_State_Summary"] == "UNASSESSED"

    readme = (SHADOW / "README.md").read_text()
    assert "$25 million tranche" in readme
    assert "title transfer" in readme
    assert "It does not establish" in readme


def test_research_queues_are_partitioned_without_invented_tasks():
    source = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
    normalized = load_workbook(MASTER_WORKBOOK, read_only=True, data_only=False)
    _, source_tasks = table(source["Contributor_Queue"])
    normalized_tasks = []
    for sheet in ("Research_Queue_T1", "Research_Queue_T2", "Research_Queue_T3"):
        _, rows = table(normalized[sheet])
        normalized_tasks.extend(rows)

    assert len(normalized_tasks) == len(source_tasks) == 13
    assert {row["Task_ID"] for row in normalized_tasks} == {row["Task_ID"] for row in source_tasks}


def test_public_contribution_forms_fail_closed_before_canonicalization():
    for path in SHADOW_ISSUE_FORMS:
        text = path.read_text()
        assert "required: true" in text
        assert "private" in text.lower()
        assert "canonical" in text.lower()

    evidence_form = SHADOW_ISSUE_FORMS[1].read_text()
    assert "Public source URL" in evidence_form
    assert "Source agency, registry, court, or publisher" in evidence_form
    assert "No public form writes directly to `CANONICAL`" in evidence_form


def test_shadow_camel_q3_release_is_immutable_and_non_authoritative():
    assert sha256(REPORT_COVER) == "a63b06e35ccb7733a464612f7d3a4e399192fc98b266467adbc3c21ab4a2230a"
    assert sha256(REPORT_PDF) == "363c2527ff153b5f5b1d2ca993ef129c5ee3a4bca2b6ee207000c96c61061000"
    assert sha256(REPORT_DOCX) == "1a5633c25fad9895eacdc29fe56cbae349c3462fc7e6e3f16d63b5d6c1ec2f0d"
    assert REPORT_PDF.read_bytes().startswith(b"%PDF-")
    assert REPORT_DOCX.read_bytes().startswith(b"PK")

    text = REPORT_PAGE.read_text()
    assert "first-edition public-research release; preliminary and non-canonical" in text
    assert "1.8 / 5 — Band 2 (Stable / Watch)" in text
    assert "not an official CAMELS rating" in text
    assert "not a `CANONICAL` evidence state" in text
    assert "not been independently reproduced" in text
    assert "No public form writes directly to `CANONICAL`" not in text
    assert "does not become `CANONICAL`" in text
