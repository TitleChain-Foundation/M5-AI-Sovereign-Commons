#!/usr/bin/env python3
"""Build the normalized public SHADOW M5Index workbook from v0.1 source files."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_MAP = {
    "Assets": "Assets_Index",
    "Ownership_Capital": "Ownership_Capital",
    "Debt_Liens": "Debt_Liens",
    "TitleChain_Asset_State": "Title_State",
    "Dispositions_Deal_Flow": "Dispositions",
    "SHADOW_Scores": "SHADOW_Scores",
    "Orbitalys_Threats": "Orbitalys_Threats",
    "Peoples_Trust_Review": "People_Trust_Review",
    "Evidence_Sources": "Evidence_Sources",
}

REQUESTED_SHEETS = [
    "README",
    "Assets",
    "Ownership_Capital",
    "Debt_Liens",
    "TitleChain_Asset_State",
    "Rights_Encumbrances",
    "Claims_Obligations",
    "Environmental_State",
    "Dispositions_Deal_Flow",
    "SHADOW_Scores",
    "SHADOW_CAMEL",
    "Orbitalys_Threats",
    "Peoples_Trust_Review",
    "Evidence_Sources",
    "Research_Queue_T1",
    "Research_Queue_T2",
    "Research_Queue_T3",
]

CAMEL_HEADERS = [
    "Institution_ID",
    "Institution_Name",
    "As_Of_Date",
    "Capital_1_5",
    "Asset_Quality_1_5",
    "Management_Risk_Proxy_1_5",
    "Earnings_1_5",
    "Liquidity_1_5",
    "Sensitivity_Proxy_1_5",
    "Composite_1_5",
    "Formula_Version",
    "Evidence_State",
    "Confidence",
    "Status",
    "Notes",
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def nonempty_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = [list(row) for row in rows[1:] if any(value not in (None, "") for value in row)]
    return headers, data


def append_table(target, headers, rows):
    target.append(headers)
    for row in rows:
        target.append(row)


def copy_sheet(source, target):
    headers, rows = nonempty_rows(source)
    append_table(target, headers, rows)


def split_records(source, record_types):
    headers, rows = nonempty_rows(source)
    type_index = headers.index("Record_Type")
    return headers, [row for row in rows if row[type_index] in record_types]


def style_sheet(ws):
    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(max((len(value) for value in values), default=0) + 2, 12), 48)
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def save_deterministic(workbook, output: Path) -> None:
    """Save an XLSX with stable metadata, member order, and ZIP timestamps."""
    fixed_time = datetime(2026, 9, 20, tzinfo=timezone.utc)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time

    with NamedTemporaryFile(suffix=".xlsx", dir=output.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        with zipfile.ZipFile(temporary, "r") as source, zipfile.ZipFile(
            output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o600 << 16
                content = source.read(name)
                if name == "docProps/core.xml":
                    content = re.sub(
                        rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                        rb"\g<1>2026-09-20T00:00:00Z\g<2>",
                        content,
                    )
                target.writestr(info, content)
    finally:
        temporary.unlink(missing_ok=True)


def build(source_workbook: Path, source_csv: Path, output: Path) -> None:
    source = load_workbook(source_workbook, data_only=False, read_only=True)
    result = Workbook()
    result.remove(result.active)

    readme = result.create_sheet("README")
    readme_rows = [
        ("SHADOW M5Index normalized public master dataset v0.1", ""),
        ("Status", "RESEARCH_DATASET — NOT CANONICAL TITLE, LEGAL, FINANCIAL, REGULATORY, OR ACQUISITION EVIDENCE"),
        ("Generated from", source_workbook.name),
        ("Source workbook SHA-256", digest(source_workbook)),
        ("Assets CSV", source_csv.name),
        ("Assets CSV SHA-256", digest(source_csv)),
        ("Normalization", "Source rows are reorganized into requested tabs; missing facts remain blank, UNASSESSED, UNRESOLVED, or NOT SCORED."),
        ("Evidence boundary", "SUBMITTED and SOURCE_VERIFIED are research workflow states. Neither means CANONICAL, title verified, closing confirmed, or approved."),
        ("Scoring boundary", "SHADOW scores are NOT SCORED. SHADOW CAMEL contains headers only and is not an official CAMELS rating."),
        ("Spring Commons boundary", "SHD-029 is an accelerated-disposition research lead. Its public-benefit conveyance path, applicant, approval, funding, and transfer remain unconfirmed."),
        ("Privacy boundary", "Public-safe research only. Do not add private M5POD, identity, KYC/KYB, bank, credential, or security-sensitive evidence."),
    ]
    for row in readme_rows:
        readme.append(row)

    for target_name, source_name in SHEET_MAP.items():
        target = result.create_sheet(target_name)
        copy_sheet(source[source_name], target)

    rights_headers, rights_rows = split_records(source["Rights_Obligations"], {"RIGHT", "ENCUMBRANCE"})
    append_table(result.create_sheet("Rights_Encumbrances"), rights_headers, rights_rows)

    claims_headers, claims_rows = split_records(source["Rights_Obligations"], {"CLAIM", "OBLIGATION"})
    append_table(result.create_sheet("Claims_Obligations"), claims_headers, claims_rows)

    env_headers, env_rows = split_records(source["Rights_Obligations"], {"ENVIRONMENTAL"})
    append_table(result.create_sheet("Environmental_State"), env_headers, env_rows)

    append_table(result.create_sheet("SHADOW_CAMEL"), CAMEL_HEADERS, [])

    asset_headers, asset_rows = nonempty_rows(source["Assets_Index"])
    asset_id_index = asset_headers.index("Asset_ID")
    tranche_index = asset_headers.index("Tranche")
    tranches = {row[asset_id_index]: row[tranche_index] for row in asset_rows}
    queue_headers, queue_rows = nonempty_rows(source["Contributor_Queue"])
    queue_asset_index = queue_headers.index("Asset_ID")
    for number in (1, 2, 3):
        tranche = f"Tranche {number}"
        rows = [row for row in queue_rows if tranches.get(row[queue_asset_index]) == tranche]
        append_table(result.create_sheet(f"Research_Queue_T{number}"), queue_headers, rows)

    result._sheets = [result[name] for name in REQUESTED_SHEETS]
    for name in REQUESTED_SHEETS:
        style_sheet(result[name])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 110
    readme.auto_filter.ref = "A1:B11"

    result.calculation.fullCalcOnLoad = True
    result.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    save_deterministic(result, output)

    with source_csv.open(newline="", encoding="utf-8-sig") as handle:
        csv_ids = [row["Asset_ID"] for row in csv.DictReader(handle)]
    output_book = load_workbook(output, read_only=True, data_only=False)
    output_ids = [row[0] for row in output_book["Assets"].iter_rows(min_row=2, values_only=True) if row[0]]
    if output_book.sheetnames != REQUESTED_SHEETS:
        raise RuntimeError(f"unexpected sheet order: {output_book.sheetnames}")
    if output_ids != csv_ids:
        raise RuntimeError("normalized Assets sheet does not match the assets CSV")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-workbook", type=Path, required=True)
    parser.add_argument("--source-assets-csv", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source_workbook, args.source_assets_csv, args.output)


if __name__ == "__main__":
    main()
